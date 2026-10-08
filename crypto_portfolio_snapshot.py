#!/usr/bin/env python3
"""
Crypto Portfolio Snapshot: a print-ready page for the binder.

Reads the Crypto Portfolio Tracker (values, costs, tiers and scores as last saved in Excel) and the Crypto Analysis
workbook (each coin's biggest risk and open point), and writes a landscape Word document (and a PDF when LibreOffice
is available). It retrieves NO new data and changes NO file.

Usage:
    python3 crypto_portfolio_snapshot.py TRACKER.xlsm [ANALYSIS.xlsx] [--outdir DIR]
With no analysis workbook named, it uses the highest-version 'Crypto Analysis V*.xlsx' it can find.
The tracker must have been opened and saved in Excel after the last price refresh (the saved values are what is printed).
"""
import sys, re, os, glob, argparse, datetime, subprocess, shutil

try:
    import openpyxl
except ImportError:
    os.system(f"{sys.executable} -m pip install openpyxl --break-system-packages -q")
    import openpyxl
try:
    from docx import Document
except ImportError:
    os.system(f"{sys.executable} -m pip install python-docx --break-system-packages -q")
    from docx import Document
from docx.shared import Pt, Inches
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

TIER_FILL = {"Exceptional": "C6E0B4", "Strong": "DDEBF7", "Average/Decent": "FFF2CC", "Weak": "F8CBAD"}
TIER_ORDER = ["Exceptional", "Strong", "Average/Decent", "Weak", "Not rated", "No rating (cash)"]


def version_key(p):
    m = re.search(r"V(\d+(?:\.\d+)*)", os.path.basename(p))
    return (tuple(int(x) for x in m.group(1).split(".")) if m else (0,), os.path.getmtime(p))


def find_analysis(arg):
    if arg:
        return arg
    c = []
    for f in ("/mnt/project", "/mnt/user-data/uploads", "."):
        c += glob.glob(os.path.join(f, "Crypto Analysis V*.xlsx"))
    if not c:
        return None
    return sorted(c, key=version_key)[-1]


def money(v, signed=False):
    if not isinstance(v, (int, float)):
        return "-"
    s = f"${abs(v):,.2f}"
    return ("-" if v < 0 else ("+" if signed and v > 0 else "")) + s


def pct(v, signed=True):
    if not isinstance(v, (int, float)):
        return "-"
    return f"{'+' if signed and v > 0 else ''}{v*100:.1f}%"


def qty(v):
    if not isinstance(v, (int, float)):
        return "-"
    return f"{v:,.8f}".rstrip("0").rstrip(".") if abs(v) < 1000 else f"{v:,.2f}"


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def setc(cell, text, size=8, bold=False, fill=None, right=False, italic=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    if right:
        p.alignment = 2
    r = p.add_run("" if text is None else str(text)); r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if fill:
        shade(cell, fill)


def grid(table, widths):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "BFBFBF"); b.append(el)
    tblPr.append(b)
    mar = OxmlElement("w:tblCellMar")
    for e, w in (("top", 15), ("bottom", 15), ("left", 60), ("right", 60)):
        el = OxmlElement(f"w:{e}"); el.set(qn("w:w"), str(w)); el.set(qn("w:type"), "dxa"); mar.append(el)
    tblPr.append(mar)
    table.autofit = False
    for row in table.rows:
        for c, w in zip(row.cells, widths):
            c.width = Inches(w)
    for gc, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 1440)))
    # keep each row on one page, and repeat the header row on every printed page
    for row in table.rows:
        cs = OxmlElement("w:cantSplit"); row._tr.get_or_add_trPr().append(cs)
    trPr = table.rows[0]._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)


def heading(doc, text, before=6):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(10)


def read_tracker(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    p = wb["Portfolio"]
    head = {str(p.cell(9, c).value).strip(): c for c in range(1, p.max_column + 1) if p.cell(9, c).value}
    need = ["Location", "Coin", "Quantity", "Total cost ($)", "Value ($)", "Unrealized P&L ($)", "Return %", "Rating tier", "Rating score"]
    miss = [h for h in need if h not in head]
    if miss:
        sys.exit(f"This tracker has no {', '.join(miss)} column(s) on Portfolio. Use Crypto Portfolio Tracker V2 or later.")
    if not isinstance(p.cell(5, 1).value, (int, float)):
        sys.exit("The tracker has no saved values. Open it in Excel, let the prices refresh, save it, then run this again.")
    rows = []
    for r in range(10, 70):
        coin = p.cell(r, head["Coin"]).value
        if not coin:
            continue
        g = lambda h: p.cell(r, head[h]).value
        q = g("Quantity")
        if isinstance(q, (int, float)) and q == 0:
            continue
        rows.append(dict(loc=g("Location"), coin=str(coin), qty=q, cost=g("Total cost ($)"), value=g("Value ($)"),
                         pl=g("Unrealized P&L ($)"), ret=g("Return %"), tier=g("Rating tier") or "", score=g("Rating score")))
    tot = dict(value=p.cell(5, 1).value, cost=p.cell(5, 3).value, pl=p.cell(5, 5).value, ret=p.cell(5, 7).value,
               missing=p.cell(5, 9).value, status=str(p.cell(6, 1).value or ""))
    rat = wb["Ratings"]
    rated_as, rsrc = {}, str(rat["A2"].value or "")
    for r in range(6, 40):
        a, b = rat.cell(r, 1).value, rat.cell(r, 2).value
        if a:
            rated_as[str(a)] = None if b == "none" else (str(b) if b else str(a))
    weighted = None
    for r in range(38, 50):
        for c in range(1, 15):
            v = rat.cell(r, c).value
            if isinstance(v, str) and v.startswith("Value-weighted score"):
                for cc in range(c + 1, 15):
                    if isinstance(rat.cell(r, cc).value, (int, float)):
                        weighted = rat.cell(r, cc).value; break
    return rows, tot, rated_as, rsrc, weighted


def read_notes(path):
    if not path:
        return {}, None
    wb = openpyxl.load_workbook(path, data_only=True)
    notes = {}
    for r in wb["Coin notes"].iter_rows(min_row=4, values_only=True):
        if r[0]:
            notes[str(r[0])] = dict(risk=r[2], open=r[3])
    return notes, os.path.basename(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("tracker")
    ap.add_argument("analysis", nargs="?")
    ap.add_argument("--outdir", default="/mnt/user-data/outputs")
    ap.add_argument("--date")
    a = ap.parse_args()
    rows, tot, rated_as, rsrc, weighted = read_tracker(a.tracker)
    apath = find_analysis(a.analysis)
    notes, aname = read_notes(apath)
    date = a.date or datetime.date.today().isoformat()
    m = re.search(r"prices as of (.+)$", tot["status"])
    prices_asof = m.group(1).strip() if m else re.sub(r"^Prices:\s*", "", tot["status"])

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Inches(11), Inches(8.5)
    sec.left_margin = sec.right_margin = Inches(0.5)
    sec.top_margin, sec.bottom_margin = Inches(0.45), Inches(0.45)
    st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(8)

    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"CRYPTO PORTFOLIO SNAPSHOT  |  {date}"); r.bold = True; r.font.size = Pt(15)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"Prices: {prices_asof}.   Holdings and costs: {os.path.basename(a.tracker)}.   "
                  f"Ratings: {aname or 'tracker Ratings tab'}.")
    r.italic = True; r.font.size = Pt(7.5)

    # totals strip
    tb = doc.add_table(rows=2, cols=5); grid(tb, [2.0] * 5)
    for i, (h, v) in enumerate((("Total value", money(tot["value"])), ("Total cost", money(tot["cost"])),
                                ("Unrealized P&L", money(tot["pl"], True)), ("Return on cost", pct(tot["ret"])),
                                ("Value-weighted rating score", f"{weighted:.1f} / 100" if isinstance(weighted, (int, float)) else "-"))):
        setc(tb.cell(0, i), h, bold=True, fill="D9D9D9")
        setc(tb.cell(1, i), v, size=11, bold=True)

    # value by tier
    by = {k: 0.0 for k in TIER_ORDER}
    for h in rows:
        t = h["tier"]
        k = "No rating (cash)" if t == "-" else ("Not rated" if (not t or t.startswith("Not rated")) else t)
        by[k] = by.get(k, 0) + (h["value"] if isinstance(h["value"], (int, float)) else 0)
    total = sum(by.values()) or 1
    heading(doc, "VALUE BY RATING TIER")
    shown = [k for k in TIER_ORDER if by.get(k)]
    tb = doc.add_table(rows=2, cols=len(shown)); grid(tb, [10.0 / len(shown)] * len(shown))
    for i, k in enumerate(shown):
        setc(tb.cell(0, i), k, bold=True, fill=TIER_FILL.get(k, "E7E6E6"))
        setc(tb.cell(1, i), f"{money(by[k])}  ({by[k]/total*100:.0f}%)", size=9)

    # holdings
    heading(doc, "HOLDINGS  (largest first)")
    rows.sort(key=lambda h: -(h["value"] if isinstance(h["value"], (int, float)) else 0))
    hdr = ["Location", "Coin", "Quantity", "Value", "Share", "Cost", "P&L", "Return", "Tier", "Score"]
    widths = [0.9, 0.7, 1.35, 1.0, 0.6, 1.0, 1.0, 0.75, 1.4, 0.6]
    tb = doc.add_table(rows=len(rows) + 1, cols=len(hdr)); grid(tb, widths)
    for i, h in enumerate(hdr):
        setc(tb.cell(0, i), h, bold=True, fill="D9D9D9", right=i in (2, 3, 4, 5, 6, 7, 9))
    for ri, h in enumerate(rows, 1):
        share = h["value"] / tot["value"] if isinstance(h["value"], (int, float)) and tot["value"] else None
        vals = [h["loc"], h["coin"], qty(h["qty"]), money(h["value"]), pct(share, False), money(h["cost"]),
                money(h["pl"], True), pct(h["ret"]), h["tier"] if h["tier"] != "-" else "no rating",
                f"{h['score']:.1f}" if isinstance(h["score"], (int, float)) else "-"]
        for i, v in enumerate(vals):
            setc(tb.cell(ri, i), v, bold=i == 1, right=i in (2, 3, 4, 5, 6, 7, 9),
                 fill=TIER_FILL.get(h["tier"]) if i == 8 else None)

    # what to watch: one line per rated coin held (by its rating coin), largest value first
    heading(doc, "WHAT TO WATCH  (from each coin's evaluation)")
    seen, watch = set(), []
    for h in rows:
        key = rated_as.get(h["coin"], h["coin"])
        if not key or key in seen:
            continue
        seen.add(key)
        held = [x["coin"] for x in rows if rated_as.get(x["coin"], x["coin"]) == key]
        n = notes.get(key, {})
        watch.append((key, sorted(set(held)), h["tier"], h["score"], n.get("risk"), n.get("open")))
    tb = doc.add_table(rows=len(watch) + 1, cols=4); grid(tb, [1.2, 1.3, 4.0, 3.5])
    for i, h in enumerate(("Coin (held as)", "Tier / score", "Biggest risk", "Open point")):
        setc(tb.cell(0, i), h, bold=True, fill="D9D9D9")
    for ri, (key, held, tier, score, risk, op) in enumerate(watch, 1):
        label = key + ("" if held == [key] else f" ({', '.join(held)})")
        setc(tb.cell(ri, 0), label, bold=True)
        setc(tb.cell(ri, 1), f"{tier}  {score:.1f}" if isinstance(score, (int, float)) else tier, fill=TIER_FILL.get(tier))
        setc(tb.cell(ri, 2), risk or "-", size=7)
        setc(tb.cell(ri, 3), op if op not in (None, "None") else "-", size=7)

    notes_line = []
    if isinstance(tot["missing"], (int, float)) and tot["missing"] > 0:
        notes_line.append(f"{int(tot['missing'])} holding(s) have no cost entered, so P&L leaves them out.")
    notes_line.append("Costs marked as estimates in the tracker's Notes column are included as entered.")
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(5)
    r = p.add_run(" ".join(notes_line) + "  Tiers: Exceptional 90–100, Strong 75–89.9, Average/Decent 60–74.9, Weak below 60. "
                  "Ratings describe the quality of retrievable facts on the date evaluated, not future price. "
                  "This is a record of holdings and an analytical rating, not financial advice.")
    r.italic = True; r.font.size = Pt(6.5)
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8)
    r = p.add_run("Notes:  ________________________________________________________________________________________________"
                  "________________________________"); r.font.size = Pt(9)

    os.makedirs(a.outdir, exist_ok=True)
    out = os.path.join(a.outdir, f"Crypto Portfolio Snapshot {date}.docx")
    doc.save(out)
    print(f"Tracker used  : {a.tracker}")
    print(f"Ratings notes : {apath or 'none found (biggest risk and open point left blank)'}")
    print(f"Saved         : {out}")
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", a.outdir, out],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
        pdf = out[:-5] + ".pdf"
        if os.path.exists(pdf):
            print(f"PDF           : {pdf}")


if __name__ == "__main__":
    main()
