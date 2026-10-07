#!/usr/bin/env python3
"""
Crypto Analysis coin report.

Reads the Crypto Analysis workbook (sheets Grades, Evidence, Coin notes, Usage method, Change log, Settings)
and writes a one-coin Word report: what it does, the result, all 15 lines with their evidence, the arithmetic,
gaps and fragility, and the coin's recent change-log entries. It retrieves NO new data and changes NO file.
The report is always regenerated from the workbook, never edited by hand.

Usage:
    python3 crypto_analysis_report.py TICKER [TICKER ...]           # newest 'Crypto Analysis V*.xlsx'
    python3 crypto_analysis_report.py TICKER --workbook FILE.xlsx --outdir /mnt/user-data/outputs
"""
import sys, re, os, glob, argparse, datetime

try:
    import openpyxl
except ImportError:
    os.system(f"{sys.executable} -m pip install openpyxl --break-system-packages -q")
    import openpyxl
try:
    from docx import Document
except ImportError:
    os.system(f"{sys.executable} -m pip install python-docx --break-system-packages -q")
    from docx import Document
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

TIER_FILL = {"Exceptional": "C6E0B4", "Strong": "DDEBF7", "Average/Decent": "FFF2CC", "Weak": "F8CBAD"}
GRADE_PCT = {"A": 1.0, "B": 0.75, "C": 0.5, "D": 0.25, "E": 0.0}


def version_key(p):
    m = re.search(r"V(\d+(?:\.\d+)*)", os.path.basename(p))
    return (tuple(int(x) for x in m.group(1).split(".")) if m else (0,), os.path.getmtime(p))


def find_workbook(arg):
    if arg:
        return arg
    cands = []
    for folder in ("/mnt/project", "/mnt/user-data/uploads", "."):
        cands += glob.glob(os.path.join(folder, "Crypto Analysis V*.xlsx"))
    if not cands:
        sys.exit("No 'Crypto Analysis V*.xlsx' found. Pass it with --workbook.")
    cands.sort(key=version_key)
    if len(cands) > 1:
        print("WARNING: several workbooks found; using the highest version:", cands[-1])
    return cands[-1]


def iso(v):
    if v is None:
        return "n/a"
    if hasattr(v, "date"):
        return v.date().isoformat()
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return str(v)


def load(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    g = wb["Grades"]
    head = [c.value for c in g[4]]
    col = {h: i for i, h in enumerate(head) if h}
    pts = [g.cell(3, col[f"L{n}"] + 1).value for n in range(1, 16)]
    coins = {}
    for r in g.iter_rows(min_row=5, values_only=True):
        if r[col["Ticker"]]:
            coins[str(r[col["Ticker"]])] = {h: r[i] for h, i in col.items()}
    names, cuts = {}, {}
    s = wb["Settings"]
    for r in s.iter_rows(values_only=True):
        if r[0] and re.fullmatch(r"L\d+", str(r[0])) and isinstance(r[2], (int, float)):
            names[int(str(r[0])[1:])] = r[1]
            cuts[int(str(r[0])[1:])] = r[3]
    thr = {}
    for r in s.iter_rows(values_only=True):
        if r[0] and isinstance(r[1], (int, float)):
            for k, key in (("Exceptional", "ex"), ("Strong", "st"), ("Average/Decent", "av"), ("Minimum coverage", "cov")):
                if str(r[0]).startswith(k):
                    thr[key] = r[1]
    ev = {}
    for r in wb["Evidence"].iter_rows(min_row=4, values_only=True):
        if r[0]:
            ev[(str(r[0]), int(str(r[1])[1:]))] = r[4]
    notes = {}
    for r in wb["Coin notes"].iter_rows(min_row=4, values_only=True):
        if r[0]:
            notes[str(r[0])] = dict(does=r[1], risk=r[2], open=r[3], caps=r[4])
    dag = set()
    for r in wb["Usage method"].iter_rows(min_row=4, values_only=True):
        if r and r[0] and len(r) > 8 and isinstance(r[8], str) and r[8].lower().startswith("yes"):
            dag.add(str(r[0]))
    log = []
    for r in wb["Change log"].iter_rows(min_row=4, values_only=True):
        if r[1]:
            log.append(r)
    return coins, pts, names, cuts, thr, ev, notes, dag, log


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def setc(cell, text, size=8, bold=False, fill=None, italic=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    r = p.add_run("" if text is None else str(text))
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if fill:
        shade(cell, fill)


def grid(table, widths):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "BFBFBF")
        b.append(el)
    tblPr.append(b)
    mar = OxmlElement("w:tblCellMar")
    for e, w in (("top", 15), ("bottom", 15), ("left", 60), ("right", 60)):
        el = OxmlElement(f"w:{e}"); el.set(qn("w:w"), str(w)); el.set(qn("w:type"), "dxa"); mar.append(el)
    tblPr.append(mar)
    table.autofit = False
    for row in table.rows:
        for c, w in zip(row.cells, widths):
            c.width = Inches(w)
    for gc, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 1440)))


def heading(doc, text, size=9.5, before=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(size)
    return p


def bullet(doc, head, text, size=8):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
    if head:
        r = p.add_run(head); r.bold = True; r.font.size = Pt(size)
    r = p.add_run(text); r.font.size = Pt(size)


def report(t, data, wbname, outdir):
    coins, pts, names, cuts, thr, ev, notes, dag, log = data
    if t not in coins:
        sys.exit(f"{t} is not on the Grades sheet.")
    c = coins[t]
    if not isinstance(c["Score"], (int, float)):
        sys.exit("No calculated scores on Grades. Recalculate the workbook (open and save it in Excel, or recalculate with LibreOffice) first.")
    n = notes.get(t, {})
    grades = [c[f"L{i}"] for i in range(1, 16)]
    tier = c["Tier"] or ""
    tier_names = {1: "Exceptional", 2: "Strong", 3: "Average/Decent", 4: "Weak"}
    score_tier = tier_names.get(c["Score tier #"], "")
    last = iso(c["Last evaluated"])

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.45)
    sec.top_margin, sec.bottom_margin = Inches(0.35), Inches(0.35)
    st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(8)

    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"CRYPTO ANALYSIS  |  {c['Name']} ({t})"); r.bold = True; r.font.size = Pt(14)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"Group: {c['Group']}   |   Last evaluated: {last}   |   Universe: spot on Coinbase, Kraken or Binance.US   |   "
                  f"Source: {wbname}")
    r.italic = True; r.font.size = Pt(7.5)

    tb = doc.add_table(rows=3, cols=2); grid(tb, [1.2, 6.3])
    for i, (k, v) in enumerate((("What it does", n.get("does")), ("Biggest risk", n.get("risk")), ("Open point", n.get("open")))):
        setc(tb.cell(i, 0), k, bold=True, fill="F2F2F2"); setc(tb.cell(i, 1), v or "")

    heading(doc, "RESULT")
    tb = doc.add_table(rows=2, cols=5); grid(tb, [1.6, 1.0, 1.4, 1.0, 2.5])
    caps = n.get("caps") or "None"
    if c["Group"] == "Memecoin" and "Memecoin" not in caps:
        caps = "Memecoin: held at Average/Decent. " + ("" if caps == "None" else caps)
    for i, h in enumerate(("Score", "Rating (÷10)", "Tier after caps", "Coverage", "Caps and flags")):
        setc(tb.cell(0, i), h, bold=True, fill="D9D9D9")
    vals = (f"{c['Score']:.1f} / 100 on {c['Scored']:g} scored pts", f"{c['Rating (/10)']:.1f} / 10", tier,
            f"{round(c['Coverage'] * 100)}%", caps)
    for i, v in enumerate(vals):
        setc(tb.cell(1, i), v, size=9 if i < 4 else 8, bold=i < 4, fill=TIER_FILL.get(tier) if i == 2 else None)

    heading(doc, "THE FIFTEEN LINES  (grade: A = 100%, B = 75%, C = 50%, D = 25%, E = 0% of the line's points; blank = not retrieved, not scored; NA = not applicable)")
    tb = doc.add_table(rows=16, cols=5); grid(tb, [1.6, 0.35, 0.45, 0.4, 4.8])
    for i, h in enumerate(("Line", "Max", "Grade", "Pts", "Key data (value | source | date)")):
        setc(tb.cell(0, i), h, bold=True, fill="D9D9D9")
    for i in range(15):
        g = grades[i] or ""
        p_ = GRADE_PCT.get(g)
        ptxt = "-" if p_ is None else f"{pts[i] * p_:g}"
        gtxt = (g or "blank") + (" †" if i == 5 and t in dag and g else "")
        evid = ev.get((t, i + 1)) or ("Not retrieved" if not g else "")
        setc(tb.cell(i + 1, 0), f"L{i+1} {names.get(i+1, '')}", size=7.5, bold=True)
        setc(tb.cell(i + 1, 1), pts[i], size=7.5)
        setc(tb.cell(i + 1, 2), gtxt, size=7.5, bold=True, fill=("F2F2F2" if not g or g == "NA" else None))
        setc(tb.cell(i + 1, 3), ptxt, size=7.5)
        setc(tb.cell(i + 1, 4), evid, size=6.5)

    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"Calculation: earned {c['Earned']:g} ÷ scored {c['Scored']:g} × 100 = {c['Score']:.1f}.  "
                  f"Applicable points {c['Applicable']:g}; coverage {c['Scored']:g} ÷ {c['Applicable']:g} = {c['Coverage']*100:.1f}%.  "
                  f"Tier from score: {score_tier}; after caps: {tier}.")
    r.font.size = Pt(8); r.bold = True

    heading(doc, "GAPS AND FRAGILITY (computed from the workbook)")
    blank = [f"L{i+1} {names.get(i+1, '')}" for i, g in enumerate(grades) if not g]
    na = [f"L{i+1}" for i, g in enumerate(grades) if g == "NA"]
    bullet(doc, "Not scored (blank, lowers coverage): ", "; ".join(blank) + "." if blank else "none.")
    if na:
        bullet(doc, "Not applicable for this coin: ", ", ".join(na) + " (removed from the denominator).")
    lines = [(thr.get("ex", 90), "Exceptional"), (thr.get("st", 75), "Strong"), (thr.get("av", 60), "Average/Decent")]
    for line, nm in lines:
        if abs(c["Score"] - line) <= 2:
            bullet(doc, "Near a tier line: ", f"score {c['Score']:.1f} is within 2 points of the {nm} line ({line}); one fact can change the tier.")
    cov_min = thr.get("cov", 0.7)
    if c["Coverage"] - cov_min <= 0.05:
        bullet(doc, "Thin coverage: ", f"{c['Coverage']*100:.0f}% is within 5 points of the {cov_min*100:.0f}% minimum for a tier.")
    if t in dag:
        bullet(doc, "Usage trend borderline (†): ", "a usage measure is within 5 points of a cut-off.")
    if blank:
        best = sum(pts[i] for i, g in enumerate(grades) if not g)
        lo = c["Earned"] / (c["Scored"] + best) * 100
        hi = (c["Earned"] + best) / (c["Scored"] + best) * 100
        bullet(doc, "If the blanks were filled: ", f"the score would fall between {lo:.1f} (all E) and {hi:.1f} (all A).")

    mine = [r for r in log if str(r[1]) == t][:4]
    heading(doc, "RECENT CHANGES (Change log, newest first; add new findings there, not here)")
    if mine:
        tb = doc.add_table(rows=len(mine) + 1, cols=5); grid(tb, [0.8, 0.9, 0.7, 3.6, 1.5])
        for i, h in enumerate(("Date", "Line", "Grade", "Value / evidence", "Source")):
            setc(tb.cell(0, i), h, bold=True, fill="D9D9D9", size=7)
        for i, r in enumerate(mine):
            old, new = r[3], r[4]
            gtxt = f"{old or '-'} → {new}" if old not in (None, "") and old != new else (new or "")
            for j, v in enumerate((iso(r[0]), r[2], gtxt, r[5], r[6])):
                setc(tb.cell(i + 1, j), v, size=6)
    else:
        bullet(doc, "", "No entries.")

    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(4)
    r = p.add_run("Tiers: Exceptional 90–100 | Strong 75–89.9 | Average/Decent 60–74.9 | Weak below 60; below 70% coverage no tier "
                  "is printed. Memecoins are held at Average/Decent. Flags lower the tier at most one step and never below "
                  "Average/Decent; Cap 4 sets Weak. This is an analytical framework, not financial advice. Ratings do not predict "
                  "price. Scores and tiers only; no buy or sell calls.")
    r.italic = True; r.font.size = Pt(6.5)

    out = os.path.join(outdir, f"Crypto Analysis Report {t} {last}.docx")
    os.makedirs(outdir, exist_ok=True)
    doc.save(out)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("tickers", nargs="+")
    ap.add_argument("--workbook")
    ap.add_argument("--outdir", default="/mnt/user-data/outputs")
    a = ap.parse_args()
    path = find_workbook(a.workbook)
    data = load(path)
    print(f"Workbook used : {path}")
    for t in a.tickers:
        print(f"Saved         : {report(t.upper(), data, os.path.basename(path), a.outdir)}")


if __name__ == "__main__":
    main()
